from pathlib import Path
from typing import Iterable
import zipfile

import pymupdf
from docx import Document as WordDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from openpyxl import load_workbook

from app.rag.types import ParsedBlock, RagProcessingError

MAX_PDF_PAGES = 1000
MAX_OFFICE_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_TEXT_CHARS = 2_000_000
MAX_PARSE_UNITS = 50_000


def _validate_office_archive(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_OFFICE_EXPANDED_BYTES:
                raise RagProcessingError("文档展开后大小超过 100 MiB 限制")
    except zipfile.BadZipFile as exc:
        raise RagProcessingError("Office 文档结构无效") from exc


def _enforce_limits(blocks: Iterable[ParsedBlock]) -> list[ParsedBlock]:
    result: list[ParsedBlock] = []
    total = 0
    for unit_count, block in enumerate(blocks, start=1):
        if unit_count > MAX_PARSE_UNITS:
            raise RagProcessingError("文档解析单元超过 50000 个限制")
        total += len(block.text)
        if total > MAX_TEXT_CHARS:
            raise RagProcessingError("文档可索引文本超过 2000000 字符限制")
        if block.text.strip():
            result.append(block)
    if not result:
        raise RagProcessingError("未提取到可索引文本，当前未启用OCR")
    return result


def parse_pdf(path: Path) -> list[ParsedBlock]:
    try:
        with pymupdf.open(path) as document:
            if document.page_count > MAX_PDF_PAGES:
                raise RagProcessingError("PDF 页数超过 1000 页限制")
            blocks = [ParsedBlock(text=page.get_text("text"), page_start=index + 1, page_end=index + 1,
                                  source_metadata={"page": index + 1})
                      for index, page in enumerate(document) if page.get_text("text").strip()]
    except RagProcessingError:
        raise
    except Exception as exc:
        raise RagProcessingError("PDF 文档解析失败") from exc
    return _enforce_limits(blocks)


def _iter_word_blocks(document: WordDocument):
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def parse_docx(path: Path) -> list[ParsedBlock]:
    _validate_office_archive(path)
    try:
        document = WordDocument(path)
        blocks: list[ParsedBlock] = []
        section: str | None = None
        table_number = 0
        for item in _iter_word_blocks(document):
            if isinstance(item, Paragraph):
                text = item.text.strip()
                if not text:
                    continue
                if item.style and item.style.name.lower().startswith("heading"):
                    section = text
                    blocks.append(ParsedBlock(text=text, section_title=section,
                                              source_metadata={"kind": "heading"}))
                else:
                    blocks.append(ParsedBlock(text=text, section_title=section,
                                              source_metadata={"kind": "paragraph"}))
            else:
                table_number += 1
                for row_number, row in enumerate(item.rows, start=1):
                    values = [cell.text.strip() for cell in row.cells]
                    if any(values):
                        blocks.append(ParsedBlock(text=" | ".join(values), section_title=section,
                                                  row_start=row_number, row_end=row_number,
                                                  source_metadata={"kind": "table", "table": table_number}))
    except RagProcessingError:
        raise
    except Exception as exc:
        raise RagProcessingError("DOCX 文档解析失败") from exc
    return _enforce_limits(blocks)


def parse_xlsx(path: Path) -> list[ParsedBlock]:
    _validate_office_archive(path)
    try:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        blocks: list[ParsedBlock] = []
        try:
            for sheet in workbook.worksheets:
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                    if values:
                        blocks.append(ParsedBlock(text=" | ".join(values), section_title=sheet.title,
                                                  row_start=row_number, row_end=row_number,
                                                  source_metadata={"kind": "worksheet", "sheet": sheet.title}))
        finally:
            workbook.close()
    except RagProcessingError:
        raise
    except Exception as exc:
        raise RagProcessingError("XLSX 文档解析失败") from exc
    return _enforce_limits(blocks)


def parse_document(path: Path, file_type: str) -> list[ParsedBlock]:
    if file_type == "PDF":
        return parse_pdf(path)
    if file_type == "DOCX":
        return parse_docx(path)
    if file_type == "XLSX":
        return parse_xlsx(path)
    raise RagProcessingError("文档类型不受知识库处理支持")
